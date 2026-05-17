from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from project_q.tools.access import resolve_allowed_path
from project_q.tools.base import ToolDefinition


SUPPORTED_SPREADSHEET_SUFFIXES = {".xlsx", ".csv", ".tsv"}


class SpreadsheetDataset:
    def __init__(self, *, path: Path, sheet_name: str, headers: list[str], rows: list[list[Any]]) -> None:
        self.path = path
        self.sheet_name = sheet_name
        self.headers = headers
        self.rows = rows


class SpreadsheetBaseTool:
    def __init__(self, workspace_root: Path, data_root: Path, settings_service=None) -> None:
        self.workspace_root = workspace_root
        self.data_root = data_root
        self.settings_service = settings_service

    def _resolve_read_path(self, raw_path: str) -> Path:
        path = resolve_allowed_path(
            raw_path,
            workspace_root=self.workspace_root,
            data_root=self.data_root,
            settings_service=self.settings_service,
            must_exist=True,
        )
        if path.suffix.lower() not in SUPPORTED_SPREADSHEET_SUFFIXES:
            raise ValueError("supported spreadsheet formats are .xlsx, .csv, and .tsv")
        return path

    def _resolve_write_path(self, raw_path: str) -> Path:
        path = resolve_allowed_path(
            raw_path,
            workspace_root=self.workspace_root,
            data_root=self.data_root,
            settings_service=self.settings_service,
        )
        if path.suffix.lower() != ".xlsx":
            raise ValueError("analysis output must be an .xlsx file")
        return path

    def _load_dataset(self, path: Path, sheet_name: str | None = None) -> SpreadsheetDataset:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            workbook = load_workbook(path, data_only=True)
            sheet = workbook[sheet_name] if sheet_name else workbook.active
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            headers, body = self._split_headers(rows)
            return SpreadsheetDataset(path=path, sheet_name=sheet.title, headers=headers, rows=body)
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = [row for row in csv.reader(handle, delimiter=delimiter)]
        headers, body = self._split_headers(rows)
        return SpreadsheetDataset(path=path, sheet_name=path.stem, headers=headers, rows=body)

    def _split_headers(self, rows: list[list[Any]]) -> tuple[list[str], list[list[Any]]]:
        for index, row in enumerate(rows[:20]):
            normalized = ["" if value is None else str(value).strip() for value in row]
            if len([value for value in normalized if value]) >= 1:
                return normalized, rows[index + 1 :]
        raise ValueError("could not find a header row in the spreadsheet")

    def _match_column(self, headers: list[str], requested_column: str | None) -> tuple[int, str]:
        if not headers:
            raise ValueError("spreadsheet has no headers")
        if not requested_column:
            numeric_headers = [index for index, header in enumerate(headers) if header]
            if not numeric_headers:
                raise ValueError("column is required")
            return numeric_headers[0], headers[numeric_headers[0]]

        requested = self._normalize(requested_column)
        best: tuple[int, str, int] | None = None
        for index, header in enumerate(headers):
            normalized = self._normalize(header)
            score = 0
            if normalized == requested:
                score = 100
            elif requested and requested in normalized:
                score = 80
            else:
                requested_terms = set(requested.split())
                header_terms = set(normalized.split())
                overlap = requested_terms & header_terms
                score = len(overlap) * 20
            if best is None or score > best[2]:
                best = (index, header, score)

        if best is None or best[2] <= 0:
            raise ValueError(f"could not match column `{requested_column}`")
        return best[0], best[1]

    def _numeric_values(self, dataset: SpreadsheetDataset, column_index: int) -> list[float]:
        values: list[float] = []
        for row in dataset.rows:
            if column_index >= len(row):
                continue
            value = row[column_index]
            numeric = self._to_number(value)
            if numeric is not None:
                values.append(numeric)
        return values

    def _calculate(self, operation: str, values: list[float]) -> float | int:
        normalized = self._normalize_operation(operation)
        if normalized == "count":
            return len(values)
        if not values:
            raise ValueError("no numeric values found for the requested column")
        if normalized == "sum":
            result = sum(values)
        elif normalized == "average":
            result = sum(values) / len(values)
        elif normalized == "min":
            result = min(values)
        elif normalized == "max":
            result = max(values)
        else:
            raise ValueError(f"unsupported operation: {operation}")
        return int(result) if float(result).is_integer() else result

    def _normalize_operation(self, operation: str) -> str:
        lowered = operation.lower().strip()
        if lowered in {"total", "sum", "add"}:
            return "sum"
        if lowered in {"avg", "average", "mean"}:
            return "average"
        if lowered in {"count", "number"}:
            return "count"
        if lowered in {"min", "minimum"}:
            return "min"
        if lowered in {"max", "maximum"}:
            return "max"
        if lowered in {"profile", "summary", "analyze", "analyse"}:
            return "profile"
        return lowered

    def _to_number(self, value: Any) -> float | None:
        if isinstance(value, bool) or value is None:
            return None
        if isinstance(value, int | float):
            return float(value)
        text = str(value).strip().replace(",", "").replace("$", "")
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None

    def _normalize(self, value: str) -> str:
        return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(value).lower())).strip()


class SpreadsheetInspectTool(SpreadsheetBaseTool):
    definition = ToolDefinition(
        tool_id="spreadsheet.inspect",
        name="Inspect Spreadsheet",
        description="Inspect workbook sheets, headers, row count, and numeric columns",
        tier=0,
    )

    def execute(self, payload: dict[str, Any]) -> dict[str, Any]:
        path = self._resolve_read_path(payload["path"])
        dataset = self._load_dataset(path, payload.get("sheet"))
        numeric_columns = []
        for index, header in enumerate(dataset.headers):
            values = self._numeric_values(dataset, index)
            if values:
                numeric_columns.append({"column": header, "numeric_count": len(values)})
        return {
            "path": str(path),
            "sheet": dataset.sheet_name,
            "headers": dataset.headers,
            "row_count": len(dataset.rows),
            "numeric_columns": numeric_columns,
            "sample_rows": dataset.rows[:5],
        }


class SpreadsheetAnalyzeTool(SpreadsheetBaseTool):
    definition = ToolDefinition(
        tool_id="spreadsheet.analyze",
        name="Analyze Spreadsheet",
        description="Calculate spreadsheet totals, averages, counts, min/max, and numeric profiles",
        tier=0,
    )

    def execute(self, payload: dict[str, Any]) -> dict[str, Any]:
        path = self._resolve_read_path(payload["path"])
        dataset = self._load_dataset(path, payload.get("sheet"))
        operation = self._normalize_operation(str(payload.get("operation", "profile")))
        if operation == "profile":
            return self._profile(dataset)

        column_index, matched_column = self._match_column(dataset.headers, payload.get("column"))
        if payload.get("group_by"):
            group_index, matched_group = self._match_column(dataset.headers, payload.get("group_by"))
            return self._grouped(dataset, operation, column_index, matched_column, group_index, matched_group)
        values = self._numeric_values(dataset, column_index)
        result = self._calculate(operation, values)
        return {
            "path": str(path),
            "sheet": dataset.sheet_name,
            "operation": operation,
            "requested_column": payload.get("column", ""),
            "matched_column": matched_column,
            "numeric_count": len(values),
            "result": result,
        }

    def _grouped(
        self,
        dataset: SpreadsheetDataset,
        operation: str,
        value_index: int,
        matched_column: str,
        group_index: int,
        matched_group: str,
    ) -> dict[str, Any]:
        grouped_values: dict[str, list[float]] = {}
        for row in dataset.rows:
            if value_index >= len(row) or group_index >= len(row):
                continue
            numeric = self._to_number(row[value_index])
            if numeric is None:
                continue
            group_key = str(row[group_index] or "(blank)").strip() or "(blank)"
            grouped_values.setdefault(group_key, []).append(numeric)
        groups = {
            key: self._calculate(operation, values)
            for key, values in sorted(grouped_values.items(), key=lambda item: item[0].lower())
        }
        return {
            "path": str(dataset.path),
            "sheet": dataset.sheet_name,
            "operation": operation,
            "requested_column": matched_column,
            "matched_column": matched_column,
            "group_by": matched_group,
            "groups": groups,
        }

    def _profile(self, dataset: SpreadsheetDataset) -> dict[str, Any]:
        columns = []
        for index, header in enumerate(dataset.headers):
            values = self._numeric_values(dataset, index)
            if not values:
                continue
            columns.append(
                {
                    "column": header,
                    "count": len(values),
                    "sum": self._calculate("sum", values),
                    "average": self._calculate("average", values),
                    "min": self._calculate("min", values),
                    "max": self._calculate("max", values),
                }
            )
        return {
            "path": str(dataset.path),
            "sheet": dataset.sheet_name,
            "operation": "profile",
            "row_count": len(dataset.rows),
            "columns": columns,
        }


class SpreadsheetWriteAnalysisTool(SpreadsheetBaseTool):
    definition = ToolDefinition(
        tool_id="spreadsheet.write_analysis",
        name="Write Spreadsheet Analysis",
        description="Create an analysis copy of a workbook with a Project Q summary sheet",
        tier=2,
    )

    def execute(self, payload: dict[str, Any]) -> dict[str, Any]:
        source_path = self._resolve_read_path(payload["path"])
        if source_path.suffix.lower() != ".xlsx":
            raise ValueError("write_analysis currently supports .xlsx workbooks")
        output_path = self._resolve_write_path(
            payload.get("output_path") or str(source_path.with_name(f"{source_path.stem}_project_q_analysis.xlsx"))
        )
        dataset = self._load_dataset(source_path, payload.get("sheet"))
        operation = self._normalize_operation(str(payload.get("operation", "profile")))
        workbook = load_workbook(source_path)
        if "Project Q Analysis" in workbook.sheetnames:
            workbook.remove(workbook["Project Q Analysis"])
        summary = workbook.create_sheet("Project Q Analysis", 0)
        summary["A1"] = "Project Q Analysis"
        summary["A3"] = "Source"
        summary["B3"] = str(source_path)
        summary["A4"] = "Result"

        if operation == "profile":
            profile = SpreadsheetAnalyzeTool(self.workspace_root, self.data_root, self.settings_service)._profile(dataset)
            summary["B4"] = "Profile"
            summary.append([])
            summary.append(["Column", "Count", "Sum", "Average", "Min", "Max"])
            for item in profile["columns"]:
                summary.append([item["column"], item["count"], item["sum"], item["average"], item["min"], item["max"]])
            result: dict[str, Any] = profile
        else:
            column_index, matched_column = self._match_column(dataset.headers, payload.get("column"))
            if payload.get("group_by"):
                group_index, matched_group = self._match_column(dataset.headers, payload.get("group_by"))
                result = SpreadsheetAnalyzeTool(self.workspace_root, self.data_root, self.settings_service)._grouped(
                    dataset,
                    operation,
                    column_index,
                    matched_column,
                    group_index,
                    matched_group,
                )
                summary["B4"] = "Grouped analysis"
                summary.append([])
                summary.append([matched_group, f"{operation} of {matched_column}"])
                for group, value in result["groups"].items():
                    summary.append([group, value])
            else:
                values = self._numeric_values(dataset, column_index)
                calculated = self._calculate(operation, values)
                summary["B4"] = calculated
                result = {
                    "operation": operation,
                    "matched_column": matched_column,
                    "numeric_count": len(values),
                    "result": calculated,
                }
            summary["A5"] = "Operation"
            summary["B5"] = operation
            summary["A6"] = "Column"
            summary["B6"] = matched_column

        for column in ("A", "B", "C", "D", "E", "F"):
            summary.column_dimensions[column].width = 22
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return {
            "source_path": str(source_path),
            "output_path": str(output_path),
            "sheet": dataset.sheet_name,
            "analysis": result,
        }
